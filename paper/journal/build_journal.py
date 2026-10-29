"""
Builds the journal-format draft (Word and LaTeX) from content_journal.py.

    python paper/journal/build_journal.py

Style parameters are collected in STYLE and are set from the journal's published format (see
journal/FORMAT_NOTES.md for the sources). Everything else is layout mechanics.
"""
from __future__ import annotations

import re
import sys
from pathlib import Path
from typing import Dict, List, Tuple

HERE = Path(__file__).resolve().parent
PAPER = HERE.parent
sys.path.insert(0, str(HERE))
sys.path.insert(0, str(PAPER))

from docx import Document  # noqa: E402
from docx.enum.section import WD_SECTION  # noqa: E402
from docx.enum.table import WD_TABLE_ALIGNMENT  # noqa: E402
from docx.enum.text import WD_ALIGN_PARAGRAPH  # noqa: E402
from docx.oxml import OxmlElement  # noqa: E402
from docx.oxml.ns import qn  # noqa: E402
from docx.shared import Cm, Pt  # noqa: E402

import content_journal as J  # noqa: E402
import figures_journal  # noqa: E402
import mathtext  # noqa: E402
from refs_journal import CITES, LONG_AUTHORS, REFERENCES  # noqa: E402

OUT = PAPER / "out"
FIGS = OUT / "journal_figs"
PIPELINE_FIG = OUT / "figs" / "pipeline.png"

STYLE = {
    "font": "Times New Roman",
    "page": (21.0, 29.7),           # A4 in cm
    "margins": (2.5, 2.5, 2.5, 2.5),  # left, right, top, bottom
    "col_width_cm": 7.7,
    "col_gap_cm": 0.6,
    "body_pt": 11,
    "title_pt": 14.3,
    "author_pt": 12,
    "abstract_pt": 10,
    "caption_pt": 10,
    "table_pt": 9,
    "ref_pt": 10,
    "table_caption_above": False,
    # Production footer of page 1. MIT Press adds this at publication; it is shown here for the demo.
    "front_lines": [
        "Transactions of the Association for Computational Linguistics, vol. 0, pp. 000–000, 2026. "
        "https://doi.org/[assigned at publication]",
        "Action Editor: [to be assigned]. Submission batch: [ ]; Revision batch: [ ]; Published [ ].",
        "© 2026 Association for Computational Linguistics. Distributed under a CC-BY 4.0 license.",
    ],
}


# Relative column widths by column count. Tables with intervals need ~1.8 cm for a value and its CI.
WIDTHS = {
    5: [2.0, 1.8, 1.0, 1.9, 0.9],
}

# ------------------------------------------------------------------ citations

CITE_RE = re.compile(r"\[@([^\]]+)\]")


class Citations:
    def __init__(self) -> None:
        self.used: List[str] = []

    def _short(self, key: str) -> Tuple[str, str]:
        if key not in CITES:
            raise KeyError(f"unknown citation key: {key}")
        if key not in self.used:
            self.used.append(key)
        return CITES[key]

    def render(self, match: re.Match) -> str:
        raw = match.group(1)
        narrative = False
        parts = [p.strip().lstrip("@") for p in raw.split(";")]
        if len(parts) == 1 and parts[0].endswith(":n"):
            narrative, parts[0] = True, parts[0][:-2]
        shorts = [self._short(k) for k in parts]
        if narrative:
            name, year = shorts[0]
            return f"{name} ({year})"
        return "(" + "; ".join(f"{n}, {y}" for n, y in shorts) + ")"


# ------------------------------------------------------------------ Word helpers

def fmt(run, size=None, bold=False, italic=False, sub=False, sup=False) -> None:
    run.font.name = STYLE["font"]
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), STYLE["font"])
    run.font.size = Pt(size or STYLE["body_pt"])
    run.font.bold = bold
    run.font.italic = italic
    run.font.subscript = sub
    run.font.superscript = sup


def para(doc, align=WD_ALIGN_PARAGRAPH.JUSTIFY, before=0.0, after=0.0, left=None, first=None, keep=False):
    p = doc.add_paragraph()
    pf = p.paragraph_format
    pf.space_before, pf.space_after = Pt(before), Pt(after)
    pf.line_spacing = 1.0
    p.alignment = align
    if left is not None:
        pf.left_indent = Cm(left)
    if first is not None:
        pf.first_line_indent = Cm(first)
    pf.keep_with_next = keep
    return p


def math_runs(p, latex: str, size=None, italic=False, bold=False) -> None:
    for text, style in mathtext.parse(latex):
        if text == "\n":
            p.add_run().add_break()
            continue
        r = p.add_run(text)
        fmt(r, size=size, bold=bold or "b" in style, italic=italic or "i" in style,
            sub="sub" in style, sup="sup" in style)


def rich(p, text: str, cites: Citations, size=None, bold=False, italic=False) -> None:
    text = (text.replace("``", "“").replace("''", "”").replace("\\_", "_")
            .replace("--", "–").replace("\\&", "&"))
    text = CITE_RE.sub(cites.render, text)
    for piece in re.split(r"(\$[^$]*\$)", text):
        if not piece:
            continue
        if piece.startswith("$") and piece.endswith("$") and len(piece) > 1:
            math_runs(p, piece[1:-1], size=size, italic=italic, bold=bold)
        else:
            fmt(p.add_run(piece), size=size, bold=bold, italic=italic)


def set_columns(section, count: int, space_cm: float) -> None:
    sect_pr = section._sectPr
    cols = sect_pr.find(qn("w:cols"))
    if cols is None:
        cols = OxmlElement("w:cols")
        sect_pr.append(cols)
    cols.set(qn("w:num"), str(count))
    cols.set(qn("w:space"), str(int(space_cm / 2.54 * 1440)))


def cell_border(cell, **edges) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.find(qn("w:tcBorders"))
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge, size in edges.items():
        el = borders.find(qn(f"w:{edge}"))
        if el is None:
            el = OxmlElement(f"w:{edge}")
            borders.append(el)
        el.set(qn("w:val"), "single" if size else "nil")
        el.set(qn("w:sz"), str(size or 0))
        el.set(qn("w:space"), "0")
        el.set(qn("w:color"), "000000")


def tight_cells(table, twips: int = 50) -> None:
    mar = OxmlElement("w:tblCellMar")
    for side in ("left", "right"):
        el = OxmlElement(f"w:{side}")
        el.set(qn("w:w"), str(twips))
        el.set(qn("w:type"), "dxa")
        mar.append(el)
    table._tbl.tblPr.append(mar)


# ------------------------------------------------------------------ document

class Counters:
    def __init__(self) -> None:
        self.eq = self.fig = self.tab = self.thm = 0


def add_caption(doc, label: str, number: int, text: str, cites: Citations, align=WD_ALIGN_PARAGRAPH.JUSTIFY):
    p = para(doc, align=align, before=3, after=8)
    fmt(p.add_run(f"{label} {number}: "), size=STYLE["caption_pt"], bold=True)
    rich(p, text, cites, size=STYLE["caption_pt"])
    return p


def add_matrix(doc, number, caption, rows, cols, values, note, cites) -> None:
    labelled = any(r for r in rows)
    ncols = len(cols) + (1 if labelled else 0)
    if STYLE["table_caption_above"]:
        add_caption(doc, "Table", number, caption, cites)
    t = doc.add_table(rows=len(rows) + 1, cols=ncols)
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    tight_cells(t)
    off = 1 if labelled else 0
    for j, c in enumerate(cols):
        cell = t.rows[0].cells[j + off]
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        fmt(p.add_run(c), size=STYLE["table_pt"], italic=True)
    for row_ in t.rows:
        for cell_ in row_.cells:
            pf = cell_.paragraphs[0].paragraph_format
            pf.space_before, pf.space_after, pf.line_spacing = Pt(1), Pt(1), 1.0
    for i, rlabel in enumerate(rows):
        if labelled:
            cell = t.rows[i + 1].cells[0]
            cell.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.LEFT
            r = cell.paragraphs[0].add_run(rlabel.replace("\\_", "_"))
            fmt(r, size=STYLE["table_pt"], italic=True)
        for j, v in enumerate(values[i]):
            cell = t.rows[i + 1].cells[j + off]
            cell.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
            fmt(cell.paragraphs[0].add_run(str(v)), size=STYLE["table_pt"])
            # square brackets: left edge of the first value column, right edge of the last
            if j == 0:
                cell_border(cell, left=8)
            if j == len(cols) - 1:
                cell_border(cell, right=8)
    if not STYLE["table_caption_above"]:
        add_caption(doc, "Table", number, caption, cites)
    if note:
        n = para(doc, after=6)
        rich(n, note, cites, size=STYLE["caption_pt"] - 1, italic=True)


def add_table(doc, number, caption, header, rows, note, cites) -> None:
    if STYLE["table_caption_above"]:
        add_caption(doc, "Table", number, caption, cites)
    t = doc.add_table(rows=1 + len(rows), cols=len(header))
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    tight_cells(t)
    total = STYLE["col_width_cm"]
    widths = WIDTHS.get(len(header))
    if widths is None:
        first = total * 0.4
        widths = [first] + [(total - first) / (len(header) - 1)] * (len(header) - 1)
    scale = total / sum(widths)
    for row in t.rows:
        for ci, cell in enumerate(row.cells):
            cell.width = Cm(widths[ci] * scale)

    def fill(cell, text, bold=False, left=False):
        p = cell.paragraphs[0]
        p.paragraph_format.space_before = Pt(1)
        p.paragraph_format.space_after = Pt(1)
        p.paragraph_format.line_spacing = 1.0
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT if left else WD_ALIGN_PARAGRAPH.CENTER
        text = text.replace("--", "–")
        m = re.match(r"^([+-]?\d+(?:\.\d+)?) \(([^)]*)\)$", text)
        if m:
            fmt(p.add_run(m.group(1)), size=STYLE["table_pt"] - 1, bold=bold)
            p.add_run().add_break()
            fmt(p.add_run(f"({m.group(2)})"), size=STYLE["table_pt"] - 1, bold=bold)
        else:
            fmt(p.add_run(text), size=STYLE["table_pt"] - 1, bold=bold)

    for ci, label in enumerate(header):
        fill(t.rows[0].cells[ci], label, bold=True, left=ci == 0)
        cell_border(t.rows[0].cells[ci], top=8, bottom=6)
    for ri, row in enumerate(rows, start=1):
        for ci, text in enumerate(row):
            fill(t.rows[ri].cells[ci], text, left=ci == 0)
            if ri == len(rows):
                cell_border(t.rows[ri].cells[ci], bottom=8)
    if not STYLE["table_caption_above"]:
        add_caption(doc, "Table", number, caption, cites)
    if note:
        n = para(doc, after=6)
        rich(n, note, cites, size=STYLE["caption_pt"] - 1, italic=True)


def add_equation(doc, latex: str, number: int) -> None:
    lines = [s.strip() for s in latex.split("\\qquad")]
    for k, line in enumerate(lines):
        p = para(doc, align=WD_ALIGN_PARAGRAPH.LEFT, before=3 if k == 0 else 0, after=3)
        stops = p.paragraph_format.tab_stops
        from docx.enum.text import WD_TAB_ALIGNMENT
        stops.add_tab_stop(Cm(STYLE["col_width_cm"] / 2), WD_TAB_ALIGNMENT.CENTER)
        stops.add_tab_stop(Cm(STYLE["col_width_cm"]), WD_TAB_ALIGNMENT.RIGHT)
        fmt(p.add_run("\t"))
        math_runs(p, line)
        if k == len(lines) - 1:
            fmt(p.add_run(f"\t({number})"))


def build_docx(path: Path) -> Citations:
    cites = Citations()
    ctr = Counters()
    doc = Document()
    st = doc.styles["Normal"]
    st.font.name = STYLE["font"]
    st.font.size = Pt(STYLE["body_pt"])

    s = doc.sections[0]
    s.page_width, s.page_height = Cm(STYLE["page"][0]), Cm(STYLE["page"][1])
    s.left_margin, s.right_margin = Cm(STYLE["margins"][0]), Cm(STYLE["margins"][1])
    s.top_margin, s.bottom_margin = Cm(STYLE["margins"][2]), Cm(STYLE["margins"][3])

    # TACL puts its production data in a footer on page 1 and a centred page number elsewhere.
    s.different_first_page_header_footer = True
    first_footer = s.first_page_footer
    first_footer.is_linked_to_previous = False
    for i, line in enumerate(STYLE["front_lines"]):
        fp = first_footer.paragraphs[0] if i == 0 else first_footer.add_paragraph()
        fp.alignment = WD_ALIGN_PARAGRAPH.LEFT
        fp.paragraph_format.space_after = Pt(0)
        fmt(fp.add_run(line), size=8)
    page_footer = s.footer
    page_footer.is_linked_to_previous = False
    pp = page_footer.paragraphs[0]
    pp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = pp.add_run()
    fmt(run, size=10)
    fld = OxmlElement("w:fldSimple")
    fld.set(qn("w:instr"), "PAGE")
    run._r.append(fld)

    # title block
    p = para(doc, align=WD_ALIGN_PARAGRAPH.CENTER, before=2, after=10)
    fmt(p.add_run(J.TITLE), size=STYLE["title_pt"], bold=True)
    tbl = doc.add_table(rows=1, cols=len(J.AUTHORS))
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    for cell, a in zip(tbl.rows[0].cells, J.AUTHORS):
        first = cell.paragraphs[0]
        first.alignment = WD_ALIGN_PARAGRAPH.CENTER
        first.paragraph_format.space_after = Pt(1)
        fmt(first.add_run(a["name"]), size=STYLE["author_pt"], bold=True)
        for line in (a["affiliation"], a["email"]):
            q = cell.add_paragraph()
            q.alignment = WD_ALIGN_PARAGRAPH.CENTER
            q.paragraph_format.space_before = Pt(0)
            q.paragraph_format.space_after = Pt(1)
            q.paragraph_format.line_spacing = 1.0
            fmt(q.add_run(line), size=STYLE["body_pt"] - 1)
    p = para(doc, after=8)

    # abstract (single column, narrower, as in the ACL family)
    p = para(doc, align=WD_ALIGN_PARAGRAPH.CENTER, before=4, after=4)
    fmt(p.add_run("Abstract"), size=STYLE["author_pt"], bold=True)
    p = para(doc, left=0.6, after=10)
    p.paragraph_format.right_indent = Cm(0.6)
    rich(p, J.ABSTRACT, cites, size=STYLE["abstract_pt"])

    body = doc.add_section(WD_SECTION.CONTINUOUS)
    set_columns(body, 2, STYLE["col_gap_cm"])

    def render_blocks(blocks) -> None:
        for block in blocks:
            kind = block[0]
            if kind == "p":
                q = para(doc, after=4, first=0.4)
                rich(q, block[1], cites)
            elif kind == "sub":
                q = para(doc, align=WD_ALIGN_PARAGRAPH.LEFT, before=5, after=2, keep=True)
                fmt(q.add_run(block[1]), bold=True)
            elif kind == "list":
                for item in block[1]:
                    q = para(doc, after=2, left=0.6, first=-0.4)
                    fmt(q.add_run("• "))
                    rich(q, item, cites)
            elif kind == "eq":
                ctr.eq += 1
                add_equation(doc, block[1], ctr.eq)
            elif kind == "fig":
                ctr.fig += 1
                fname, caption, width = block[1], block[2], block[3]
                q = para(doc, align=WD_ALIGN_PARAGRAPH.CENTER, before=4, after=2, keep=True)
                src = PIPELINE_FIG if fname == "pipeline.png" else FIGS / fname
                full = width >= 0.95
                q.add_run().add_picture(str(src), width=Cm((STYLE["col_width_cm"] * 2 + STYLE["col_gap_cm"]
                                                            if full else STYLE["col_width_cm"] * min(width * 1.7, 1.0))))
                add_caption(doc, "Figure", ctr.fig, caption, cites)
            elif kind == "table":
                ctr.tab += 1
                add_table(doc, ctr.tab, block[1], block[2], block[3], block[4], cites)
            elif kind == "matrix":
                ctr.tab += 1
                add_matrix(doc, ctr.tab, block[1], block[2], block[3], block[4], block[5], cites)
            elif kind == "thm":
                ctr.thm += 1
                q = para(doc, before=4, after=4)
                fmt(q.add_run(f"{block[1]} {ctr.thm} ({block[2].replace('$', '')}). "), bold=True)
                rich(q, block[3], cites, italic=True)
            elif kind == "proof":
                for n, text in enumerate(block[1]):
                    q = para(doc, after=3, first=0.4 if n else 0)
                    if n == 0:
                        fmt(q.add_run("Proof. "), italic=True)
                    rich(q, text, cites)
                    if n == len(block[1]) - 1:
                        fmt(q.add_run(" ∎"))
            elif kind == "remark":
                q = para(doc, after=4)
                fmt(q.add_run("Remark. "), italic=True)
                rich(q, block[1], cites)

    for number, (title, blocks) in enumerate(J.SECTIONS, start=1):
        h = para(doc, align=WD_ALIGN_PARAGRAPH.LEFT, before=9, after=3, keep=True)
        fmt(h.add_run(f"{number} {title}"), size=STYLE["body_pt"] + 1, bold=True)
        render_blocks([b if b[0] != "sub" else ("sub", f"{number}.{_sub_index(blocks, b)} {b[1]}") for b in blocks])

    # Limitations and Acknowledgments are unnumbered, between the Conclusion and the References.
    for title, blocks in J.LIMITATIONS:
        h = para(doc, align=WD_ALIGN_PARAGRAPH.LEFT, before=9, after=3, keep=True)
        fmt(h.add_run(title), size=STYLE["body_pt"] + 1, bold=True)
        render_blocks(blocks)

    h = para(doc, align=WD_ALIGN_PARAGRAPH.LEFT, before=9, after=3, keep=True)
    fmt(h.add_run("Acknowledgments"), size=STYLE["body_pt"] + 1, bold=True)
    q = para(doc, after=4)
    rich(q, J.ACKNOWLEDGMENTS, cites)

    # references
    h = para(doc, align=WD_ALIGN_PARAGRAPH.LEFT, before=9, after=3, keep=True)
    fmt(h.add_run("References"), size=STYLE["body_pt"] + 1, bold=True)
    for key in sorted(cites.used, key=lambda k: REFERENCES[k]["sort"]):
        q = para(doc, align=WD_ALIGN_PARAGRAPH.LEFT, after=3, left=0.4, first=-0.4)
        _reference(q, REFERENCES[key])

    # appendices
    for letter, (title, blocks) in zip("AB", J.APPENDICES):
        h = para(doc, align=WD_ALIGN_PARAGRAPH.LEFT, before=9, after=3, keep=True)
        fmt(h.add_run(f"{letter} {title}"), size=STYLE["body_pt"] + 1, bold=True)
        render_blocks(blocks)

    path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(path))
    return cites


def _sub_index(blocks, target) -> int:
    n = 0
    for b in blocks:
        if b[0] == "sub":
            n += 1
            if b is target:
                return n
    return n


def _reference(p, ref: Dict[str, str]) -> None:
    """Reference entry: plain text with the venue or title in italics where the entry marks it."""
    text = ref["text"]
    parts = re.split(r"(<i>.*?</i>)", text)
    for part in parts:
        if part.startswith("<i>"):
            fmt(p.add_run(part[3:-4]), size=STYLE["ref_pt"], italic=True)
        elif part:
            fmt(p.add_run(part), size=STYLE["ref_pt"])


# ------------------------------------------------------------------ LaTeX (official TACL template)

def tex_escape_text(text: str) -> str:
    """Escapes plain-text segments; math segments are passed through by the caller."""
    text = text.replace("\\_", "\u0001")
    for ch, rep in (("%", "\\%"), ("&", "\\&"), ("#", "\\#"), ("_", "\\_")):
        text = re.sub(r"(?<!\\)" + re.escape(ch), lambda _m, r=rep: r, text)
    return text.replace("\u0001", "\\_")


def tex_cites(text: str, used: set) -> str:
    """[@key] -> \\citep{key}; [@key:n] -> \\citet{key}; [@a; @b] -> \\citep{a,b} (natbib)."""
    def repl(m: re.Match) -> str:
        parts = [q.strip().lstrip("@") for q in m.group(1).split(";")]
        narrative = len(parts) == 1 and parts[0].endswith(":n")
        keys = [q[:-2] if q.endswith(":n") else q for q in parts]
        for k in keys:
            if k not in CITES:
                raise KeyError(f"unknown citation key: {k}")
            used.add(k)
        return ("\\citet{" if narrative else "\\citep{") + ",".join(keys) + "}"
    return CITE_RE.sub(repl, text)


def tex_rich(text: str, used: set) -> str:
    text = text.replace("\\&", "&")
    text = tex_cites(text, used)
    out = []
    for piece in re.split(r"(\$[^$]*\$|\\cite[pt]\{[^}]*\})", text):
        if (piece.startswith("$") and piece.endswith("$") and len(piece) > 1) or piece.startswith("\\cite"):
            out.append(piece)
        else:
            out.append(tex_escape_text(piece))
    return "".join(out)


def tex_matrix(rows, cols, values) -> str:
    labelled = any(r for r in rows)
    spec = ("l" if labelled else "") + "c" * len(cols)
    lines = [f"\\[\\begin{{array}}{{{spec}}}"]
    head = " & ".join(([""] if labelled else []) + [f"\\mathit{{{c}}}" if c else "" for c in cols])
    lines.append(head + " \\\\")
    for i, r in enumerate(rows):
        body = " & ".join(str(v) for v in values[i])
        label = f"\\mathit{{{r}}} & " if labelled else ""
        lines.append(f"{label}\\left[\\begin{{matrix}} {body} \\end{{matrix}}\\right] \\\\")
    lines.append("\\end{array}\\]")
    return "\n".join(lines)


def build_tex(path: Path, used: set) -> None:
    """
    Source in the structure of the official template (tacl2021v1-template.tex). Put tacl2021v1.sty
    and acl_natbib.bst next to it (https://transacl.org/tacl-submission-templates/). The entries in
    thebibliography are written out so no .bib file is needed. This file has not been compiled here.
    """
    head: List[str] = [
        "% TACL draft generated by paper/journal/build_journal.py.",
        "% Needs tacl2021v1.sty from https://transacl.org/tacl-submission-templates/ in the same folder.",
        "\\documentclass[11pt,a4paper]{article}",
        "\\usepackage{times,latexsym}",
        "\\usepackage{url}",
        "\\usepackage[T1]{fontenc}",
        "\\usepackage[]{tacl2021v1}",
        "\\usepackage{xspace,mfirstuc,tabulary}",
        "\\usepackage{amsmath,amssymb,amsthm,graphicx,booktabs}",
        "\\newtheorem{theorem}{Theorem}",
        "\\newtheorem{proposition}[theorem]{Proposition}",
        "\\newtheorem{remark}{Remark}",
        "",
        "\\title{" + tex_rich(J.TITLE, used) + "}",
        "\\author{" + " \\\\And ".join(
            # braces: a bare "[...]" after \\ would be read as a line-spacing argument
            f"{a['name']} \\\\ {a['affiliation']} \\\\ {{{a['email']}}}" for a in J.AUTHORS) + "}",
        "\\date{}",
        "\\begin{document}",
        "\\maketitle",
        "\\begin{abstract}",
        tex_rich(J.ABSTRACT, used),
        "\\end{abstract}",
        "",
    ]
    ctr = Counters()

    def blocks_to_tex(out: List[str], blocks) -> None:
        for b in blocks:
            kind = b[0]
            if kind == "p":
                out.extend([tex_rich(b[1], used), ""])
            elif kind == "sub":
                title = re.sub(r"^[A-Z]?\d+(\.\d+)?\s+", "", b[1])
                out.extend([f"\\subsection{{{tex_rich(title, used)}}}", ""])
            elif kind == "list":
                out.extend(["\\begin{itemize}"] + [f"\\item {tex_rich(i, used)}" for i in b[1]] + ["\\end{itemize}", ""])
            elif kind == "eq":
                parts = [q.strip() for q in b[1].split("\\qquad")]
                if len(parts) == 1:
                    out.extend(["\\begin{equation}", parts[0], "\\end{equation}", ""])
                else:
                    out.extend(["\\begin{equation}", "\\begin{aligned}", " \\\\\n".join(parts), "\\end{aligned}",
                                "\\end{equation}", ""])
            elif kind == "fig":
                ctr.fig += 1
                src = "figs/pipeline.png" if b[1] == "pipeline.png" else f"journal_figs/{b[1]}"
                out.extend(["\\begin{figure}[t]", "\\centering", f"\\includegraphics[width=\\columnwidth]{{{src}}}",
                            f"\\caption{{{tex_rich(b[2], used)}}}", "\\end{figure}", ""])
            elif kind == "table":
                header, rows = b[2], b[3]
                spec = "l" + "c" * (len(header) - 1)
                out.extend(["\\begin{table}[t]", "\\centering\\footnotesize", f"\\begin{{tabular}}{{{spec}}}", "\\toprule",
                            " & ".join(tex_escape_text(h) for h in header) + " \\\\", "\\midrule"])
                for r in rows:
                    out.append(" & ".join(tex_escape_text(c) for c in r) + " \\\\")
                out.extend(["\\bottomrule", "\\end{tabular}", f"\\caption{{{tex_rich(b[1], used)}}}"])
                if b[4]:
                    out.append(f"\\par{{\\scriptsize {tex_rich(b[4], used)}}}")
                out.extend(["\\end{table}", ""])
            elif kind == "matrix":
                out.extend(["\\begin{table}[t]", "\\centering", tex_matrix(b[2], b[3], b[4]),
                            f"\\caption{{{tex_rich(b[1], used)}}}"])
                if b[5]:
                    out.append(f"\\par{{\\scriptsize {tex_rich(b[5], used)}}}")
                out.extend(["\\end{table}", ""])
            elif kind == "thm":
                env = "theorem" if b[1] == "Theorem" else "proposition"
                out.extend([f"\\begin{{{env}}}[{tex_rich(b[2], used)}]", tex_rich(b[3], used), f"\\end{{{env}}}", ""])
            elif kind == "proof":
                out.append("\\begin{proof}")
                out.append("\n\n".join(tex_rich(t, used) for t in b[1]))
                out.extend(["\\end{proof}", ""])
            elif kind == "remark":
                out.extend(["\\begin{remark}", tex_rich(b[1], used), "\\end{remark}", ""])

    body: List[str] = []
    for title, blocks in J.SECTIONS:
        body.extend([f"\\section{{{tex_rich(title, used)}}}", ""])
        blocks_to_tex(body, blocks)
    for title, blocks in J.LIMITATIONS:
        body.extend([f"\\section*{{{title}}}", ""])
        blocks_to_tex(body, blocks)
    body.extend(["\\section*{Acknowledgments}", tex_rich(J.ACKNOWLEDGMENTS, used), ""])

    appendix: List[str] = ["\\appendix"]
    for title, blocks in J.APPENDICES:
        appendix.extend([f"\\section{{{tex_rich(title, used)}}}", ""])
        blocks_to_tex(appendix, blocks)

    # Bibliography is written last, once every citation (including appendix ones) has been seen.
    bib: List[str] = ["\\begin{thebibliography}{99}"]
    for key in sorted(used, key=lambda k: REFERENCES[k]["sort"]):
        name, year = CITES[key]
        label = f"{name.replace(' et al.', ' et~al.')}({year}){LONG_AUTHORS[key]}"
        parts = re.split(r"(<i>.*?</i>)", REFERENCES[key]["text"])
        rendered = "".join(
            ("\\textit{" + tex_escape_text(q[3:-4]) + "}") if q.startswith("<i>") else tex_escape_text(q)
            for q in parts if q
        )
        bib.extend([f"\\bibitem[{label}]{{{key}}}", rendered])
    bib.extend(["\\end{thebibliography}", ""])

    text = "\n".join(head + body + bib + appendix + ["\\end{document}", ""])
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(text, encoding="utf-8")


def main() -> None:
    OUT.mkdir(parents=True, exist_ok=True)
    figures_journal.build_all(FIGS)
    cites = build_docx(OUT / "SyntheticSilence_Journal_Draft.docx")
    tex_used: set = set()
    build_tex(OUT / "SyntheticSilence_Journal_Draft.tex", tex_used)
    assert tex_used == set(cites.used), "Word and LaTeX outputs cite different references"
    print("cited references:", len(cites.used))
    print("wrote", OUT / "SyntheticSilence_Journal_Draft.docx")
    print("wrote", OUT / "SyntheticSilence_Journal_Draft.tex")


if __name__ == "__main__":
    main()
